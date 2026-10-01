"""Editable Word resume exports from saved data; no AI or browser required."""
from html.parser import HTMLParser
from io import BytesIO
import re

from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from app.schemas.models import ResumeData, SectionMeta

TEMPLATES = {"swiss-single", "swiss-two-column", "modern", "modern-two-column", "latex", "clean", "vivid"}
ACCENTS = {"blue": "1D4ED8", "green": "15803D", "orange": "EA580C", "red": "DC2626"}

DEFAULT_SECTIONS = [
    ("personalInfo", "Personal Info", "personalInfo"),
    ("summary", "Summary", "text"),
    ("workExperience", "Experience", "itemList"),
    ("education", "Education", "itemList"),
    ("personalProjects", "Projects", "itemList"),
    ("additional", "Skills & Awards", "stringList"),
]


class _Text(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.suppressed = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.suppressed += 1
        elif not self.suppressed and tag in {"br", "p", "li", "div"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.suppressed = max(0, self.suppressed - 1)
        elif not self.suppressed and tag in {"p", "li", "div"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.suppressed:
            self.parts.append(data)


def plain_text(value: str | None) -> str:
    """Strip editor markup and XML-invalid controls without executing content."""
    value = value or ""
    if re.search(r"</?(?:p|div|br|strong|em|b|i|u|a|ul|ol|li|span|script|style)\b", value, re.I):
        parser = _Text()
        parser.feed(value)
        value = "".join(parser.parts)
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", value)
    return "\n".join(line.strip() for line in value.splitlines() if line.strip())


def render_resume_docx(
    data: ResumeData, *, template: str = "swiss-single", accent_color: str = "blue", page_size: str = "A4",
    margins: tuple[int, int, int, int] = (10, 10, 10, 10),
    section_spacing: int = 3, item_spacing: int = 2, line_height: int = 3,
    font_size: int = 3, header_scale: int = 3,
    header_font: str = "serif", body_font: str = "sans-serif", compact: bool = False,
) -> bytes:
    """Editable template layouts using the preview's typography and spacing."""
    if template not in TEMPLATES or accent_color not in ACCENTS:
        raise ValueError("Unsupported Word template or accent color")
    two_column = template in {"swiss-two-column", "modern-two-column", "vivid"}
    colored = template in {"modern", "modern-two-column", "vivid"}
    accent = ACCENTS[accent_color]
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5) if page_size == "LETTER" else Mm(210)
    section.page_height = Inches(11) if page_size == "LETTER" else Mm(297)
    section.top_margin, section.bottom_margin, section.left_margin, section.right_margin = map(Mm, margins)
    width = section.page_width - section.left_margin - section.right_margin
    # CSS pixels are 3/4 of a Word point. Families match the Windows browser's
    # serif, system sans, and monospace choices; Word can substitute if absent.
    base = (11, 12, 14, 15, 16)[font_size - 1] * .75
    multiplier = .6 if compact else 1
    item_gap = (2, 4, 8, 12, 16)[item_spacing - 1] * .75 * multiplier
    section_gap = (6, 10, 16, 20, 24)[section_spacing - 1] * .75 * multiplier
    line = (1.15, 1.25, 1.35, 1.45, 1.55)[line_height - 1] * (.92 if compact else 1)
    families = {"serif": "Georgia", "sans-serif": "Segoe UI", "mono": "Consolas"}

    def precise_size(target, size):
        # OOXML sizes use half points. Preserve fractional CSS widths with
        # Word's native text scale instead of truncating every smaller run.
        rounded = round(size * 2) / 2
        target.font.size = Pt(rounded)
        props = target.element.get_or_add_rPr() if hasattr(target, "element") else target._r.get_or_add_rPr()
        scale = props.find(qn("w:w"))
        if scale is None:
            scale = OxmlElement("w:w")
            props.append(scale)
        scale.set(qn("w:val"), str(round(size / rounded * 100)))

    def font(style, family, size, bold=False, color="1F2937"):
        style.font.name = family
        precise_size(style, size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        # Remove theme overrides inherited from python-docx's template.
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + attr), None)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = Pt(size * line)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)

    font(doc.styles["Normal"], families[body_font], base)
    font(doc.styles["List Bullet"], families[body_font], base * .92)
    font(doc.styles["Title"], families[header_font], base * (1.5, 1.75, 2, 2.25, 2.5)[header_scale - 1], True, "000000")
    font(doc.styles["Heading 1"], families[header_font], base * (1, 1.1, 1.2, 1.3, 1.4)[header_scale - 1], True, "000000")
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(section_gap)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(item_gap)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = ""

    if colored:
        doc.styles["Heading 1"].font.color.rgb = RGBColor.from_string(accent)
    if template == "clean":
        doc.styles["Title"].font.bold = False
        doc.styles["Heading 1"].font.color.rgb = RGBColor.from_string("6B7280")
        precise_size(doc.styles["Heading 1"], base * (1, 1.1, 1.2, 1.3, 1.4)[header_scale - 1] * 1.15)
    if template == "vivid":
        doc.styles["Title"].font.color.rgb = RGBColor.from_string(accent)
        doc.styles["Heading 1"].font.small_caps = True
    container = doc
    vivid_number = None
    if template == "vivid":
        numbering = doc.part.numbering_part.element
        abstract_id = max(int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))) + 1
        number_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, value in (("start", "1"), ("numFmt", "bullet"), ("lvlText", "▸"), ("lvlJc", "left")):
            el = OxmlElement("w:" + tag)
            el.set(qn("w:val"), value)
            level.append(el)
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), accent)
        props.append(color)
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Segoe UI")
        fonts.set(qn("w:hAnsi"), "Segoe UI")
        props.append(fonts)
        level.append(props)
        abstract.append(level)
        numbering.append(abstract)
        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(number_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abstract_id))
        number.append(ref)
        numbering.append(number)
        vivid_number = number_id

    def rule(p, size="6", color="9CA3AF"):
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in {"val": "single", "sz": size, "space": "2", "color": color}.items():
            bottom.set(qn("w:" + key), value)
        borders.append(bottom)
        p._p.get_or_add_pPr().append(borders)

    def paragraph(text, *, bold=None, style=None):
        cleaned = plain_text(text)
        if not cleaned:
            return
        if style == "Heading 1" and template not in {"latex", "vivid"}:
            cleaned = cleaned.upper()
        p = container.add_paragraph(style=style)
        p.add_run(cleaned).bold = bold
        if style == "List Bullet" and vivid_number is not None:
            props = p._p.get_or_add_pPr().get_or_add_numPr()
            props.get_or_add_ilvl().val = 0
            props.get_or_add_numId().val = vivid_number
        if style == "Heading 1" and template != "vivid":
            rule(p, size="12" if colored else "6", color=accent if colored else "000000" if template == "latex" else "9CA3AF")
        return p

    def small(p, *, mono=False, factor=.92):
        if p is not None:
            for run in p.runs:
                run.font.name = "Consolas" if mono else families[body_font]
                precise_size(run, base * factor)
            # Description rows retain the preview body line box.
            p.paragraph_format.line_spacing = Pt(base * factor * line)
            if p.style.name == "List Bullet":
                p.paragraph_format.line_spacing = Pt(base * line)

    def row(left, right, *, bold=False, subtitle=False):
        p = paragraph(left, bold=bold)
        if p is None:
            p = container.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        if right:
            run = p.add_run("\t" + plain_text(right))
            run.bold = False
            run.font.name = families[body_font] if subtitle else "Consolas"
            precise_size(run, base * (.95 if subtitle else .75))
            run.font.color.rgb = RGBColor.from_string("374151" if subtitle else "4B5563")
        p.paragraph_format.space_after = Pt(item_gap if subtitle else item_gap * .6)
        if subtitle:
            small(p, factor=.95)
            for run in p.runs:
                if body_font == "sans-serif":
                    run.font.name = "Segoe UI Semibold"
                else:
                    run.bold = True
        p.paragraph_format.keep_with_next = True
        return p

    def items(entries):
        for index, item in enumerate(entries):
            if getattr(item, "institution", None) is not None:
                title, subtitle = item.institution, item.degree
            else:
                title = getattr(item, "title", None) or getattr(item, "name", "")
                subtitle = getattr(item, "company", None) or getattr(item, "role", None) or getattr(item, "subtitle", None)
            first = len(container.paragraphs)
            if template in {"latex", "clean", "vivid"} and getattr(item, "company", None):
                title, subtitle = subtitle, title
            if template in {"clean", "vivid"} and subtitle:
                title = title + " | " + subtitle
                subtitle = None
            p = row(title, item.years, bold=True)
            # Word takes max(before, after); CSS adds the item and bullet gaps.
            # Put the complete entry gap on this row so it cannot collapse.
            p.paragraph_format.space_before = Pt(item_gap * 2 if index else 0)
            if subtitle or getattr(item, "location", None):
                secondary = row(subtitle, getattr(item, "location", None), subtitle=True)
                if template == "latex":
                    for run in secondary.runs:
                        run.bold = False
                        run.italic = True
            for key in ("github", "website"):
                p = paragraph(getattr(item, key, None))
                small(p, mono=True, factor=.75)
            description = item.description
            if isinstance(description, str):
                small(paragraph(description))
            else:
                styles = getattr(item, "descriptionStyles", [])
                for index, text in enumerate(description or []):
                    plain = index < len(styles) and styles[index] == "plain"
                    p = paragraph(text, style=None if plain else "List Bullet")
                    if p is not None:
                        small(p)
                        # Compact serif text uses the smaller CSS line box;
                        # keeping the full body box stretches every bullet.
                        description_scale = .92 if compact and body_font == "serif" else 1
                        p.paragraph_format.line_spacing = Pt(base * description_scale * line)
                        p.paragraph_format.space_after = Pt(item_gap * .75)
                        if not plain:
                            p.paragraph_format.left_indent = Pt(18)
                            p.paragraph_format.first_line_indent = Pt(-9)
            # Keep an entry's title with its content, while allowing long
            # description lists to flow naturally across pages.
            if len(container.paragraphs) > first:
                container.paragraphs[-1].paragraph_format.keep_with_next = False

    metadata = data.sectionMeta or [SectionMeta(id=key, key=key, displayName=label, sectionType=kind, order=index) for index, (key, label, kind) in enumerate(DEFAULT_SECTIONS)]
    visible = sorted((s for s in metadata if s.isVisible), key=lambda s: s.order)
    if any(s.key == "personalInfo" for s in visible):
        info = data.personalInfo
        p = paragraph(info.name if template in {"latex", "clean", "vivid"} else info.name.upper(), style="Title")
        if p is not None:
            if template == "vivid" and " " in info.name:
                first, rest = plain_text(info.name).split(" ", 1)
                p.runs[0].text = first + " "
                run = p.add_run(rest)
                run.bold = False
                light = "".join(f"{round(int(accent[i:i+2], 16) * .6 + 255 * .4):02X}" for i in (0, 2, 4))
                run.font.color.rgb = RGBColor.from_string(light)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
        p = paragraph(info.title if template in {"latex", "clean"} else info.title.upper())
        if p is not None:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            small(p, mono=True, factor=.82)
            p.paragraph_format.space_after = Pt(9)
            p.paragraph_format.keep_with_next = True
        contacts = [(info.email, "mailto:"), (info.phone, "tel:"), (info.location, ""),
                    (info.website, "https://"), (info.linkedin, "https://"), (info.github, "https://")]
        if any(value for value, _ in contacts):
            p = container.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            used = 0
            max_chars = (width / 12700) / (base * .82 * .6)
            for value, prefix in contacts:
                if not value:
                    continue
                value = plain_text(value)
                if used:
                    if used + len(value) + 2 > max_chars:
                        p.add_run("\n")
                        used = 0
                    else:
                        p.add_run(", ")
                        used += 2
                if prefix:
                    target = value if value.startswith(("https://", "http://")) else prefix + value
                    link = OxmlElement("w:hyperlink")
                    link.set(qn("r:id"), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
                    run = OxmlElement("w:r")
                    props = OxmlElement("w:rPr")
                    fonts = OxmlElement("w:rFonts")
                    fonts.set(qn("w:ascii"), "Consolas")
                    fonts.set(qn("w:hAnsi"), "Consolas")
                    props.append(fonts)
                    for tag, val in (("sz", str(round(base * .82 * 2))), ("color", "4B5563"), ("u", "single")):
                        element = OxmlElement("w:" + tag)
                        element.set(qn("w:val"), val)
                        props.append(element)
                    run.append(props)
                    text = OxmlElement("w:t")
                    text.text = value
                    run.append(text)
                    link.append(run)
                    p._p.append(link)
                else:
                    p.add_run(value)
                used += len(value)
            small(p, mono=True, factor=.82)
            # A different border width prevents Word merging this paragraph's
            # rule with the next section heading's border.
            if template not in {"latex", "clean", "vivid"}:
                rule(p, size="5", color=accent if colored else "9CA3AF")
            p.paragraph_format.space_after = Pt(item_gap * 1.25)
            p.paragraph_format.keep_with_next = True
    columns = None
    if two_column:
        ratio = .65 if template == "modern-two-column" else .63 if template == "vivid" else .62
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = int(width * ratio)
        table.columns[1].width = int(width * (1 - ratio))
        columns = table.rows[0].cells
        for i, cell in enumerate(columns):
            cell.width = table.columns[i].width
            props = cell._tc.get_or_add_tcPr()
            padding = OxmlElement("w:tcMar")
            for side, val in (("top", 0), ("bottom", 0), ("left", 0 if i == 0 else 160), ("right", 160 if i == 0 else 0)):
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), str(val))
                el.set(qn("w:type"), "dxa")
                padding.append(el)
            props.append(padding)
            cell._tc.remove(cell.paragraphs[0]._p)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + side)
            el.set(qn("w:val"), "single" if side == "insideV" and template == "modern-two-column" else "nil")
            if side == "insideV" and template == "modern-two-column":
                el.set(qn("w:color"), accent)
                el.set(qn("w:sz"), "12")
            borders.append(el)
        table._tbl.tblPr.append(borders)
    for meta in visible:
        if meta.key == "personalInfo":
            continue
        if columns is not None:
            sidebar = meta.key in {"education", "additional"}
            container = columns[1 if sidebar else 0]
            width = container.width - Pt(8)
        if meta.isDefault:
            value = getattr(data, meta.key, None)
            if meta.key == "additional":
                if not any(data.additional.model_dump().values()):
                    continue
                if not two_column:
                    paragraph(meta.displayName, style="Heading 1")
                labels = {"technicalSkills": "Technical Skills", "languages": "Languages", "certificationsTraining": "Certifications and Training", "awards": "Awards"}
                for key, label in labels.items():
                    values = getattr(data.additional, key)
                    if values:
                        if columns is not None:
                            container = columns[0 if key == "certificationsTraining" else 1]
                            paragraph(label, style="Heading 1")
                        p = container.add_paragraph()
                        p.paragraph_format.left_indent = Pt(0 if two_column else 96)
                        p.paragraph_format.first_line_indent = Pt(0 if two_column else -96)
                        if not two_column:
                            p.paragraph_format.tab_stops.add_tab_stop(Pt(96))
                        if not two_column:
                            p.add_run(label + ":\t").bold = True
                        p.add_run(", ".join(plain_text(v) for v in values))
                        small(p)
                        # Compact serif text uses the smaller CSS line box;
                        # keeping the full body box stretches every bullet.
                        description_scale = .92 if compact and body_font == "serif" else 1
                        p.paragraph_format.line_spacing = Pt(base * description_scale * line)
                        p.paragraph_format.space_after = Pt(item_gap * .6)
            elif value:
                paragraph(meta.displayName, style="Heading 1")
                if isinstance(value, str):
                    p = paragraph(value)
                    if p is not None:
                        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                elif isinstance(value, list):
                    items(value)
        else:
            custom = data.customSections.get(meta.key)
            if custom is None or not (custom.text or custom.items or custom.strings):
                continue
            paragraph(meta.displayName, style="Heading 1")
            if custom.text:
                paragraph(custom.text)
            if custom.items:
                items(custom.items)
            for text in custom.strings or []:
                paragraph(text, style="List Bullet")
    if columns is not None:
        for cell in columns:
            if not cell.paragraphs:
                cell.add_paragraph()
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
