"""Word content fidelity and real HTTP owner authorization."""
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from httpx import ASGITransport, AsyncClient
import pytest

from app.auth import current_user_id
from app.config import settings
from app.main import app
from app.models import User
from app.passwords import hash_password
from app.schemas.models import ResumeData
from app.word_export import render_resume_docx


def test_word_export_preserves_content_order_visibility_and_editable_text(sample_resume):
    sample_resume['personalInfo']['name'] = 'José Example'
    sample_resume['summary'] = '<p>Python <strong>developer</strong> &amp; student</p>'
    sample_resume['sectionMeta'] = [
        dict(id='personalInfo', key='personalInfo', displayName='Contact', sectionType='personalInfo', order=0),
        dict(id='custom_1', key='custom_1', displayName='Leadership', sectionType='itemList', isDefault=False, order=1),
        dict(id='summary', key='summary', displayName='Profile', sectionType='text', order=2),
        dict(id='workExperience', key='workExperience', displayName='Experience', sectionType='itemList', order=3, isVisible=False),
        dict(id='education', key='education', displayName='Education', sectionType='itemList', order=4),
        dict(id='additional', key='additional', displayName='Skills', sectionType='stringList', order=5),
    ]
    sample_resume['customSections'] = {'custom_1': {'sectionType':'itemList', 'items':[{'title':'Volunteer', 'subtitle':'Example Club', 'years':'2026', 'description':['<p>Built <em>beds</em></p>', 'Plain point'], 'descriptionStyles':['bullet','plain']}]}}
    content = render_resume_docx(ResumeData.model_validate(sample_resume), page_size="LETTER")
    with ZipFile(BytesIO(content)) as archive:
        assert 'word/document.xml' in archive.namelist()
        assert not any('media/' in name for name in archive.namelist())
    document = Document(BytesIO(content))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'JOSÉ EXAMPLE' in text
    assert 'Python developer & student' in text
    assert 'Built beds' in text and 'Plain point' in text
    assert text.index('LEADERSHIP') < text.index('PROFILE') < text.index('EDUCATION')
    assert 'Acme Corp' not in text
    assert '<strong>' not in text
    assert document.sections[0].page_width.inches == 8.5
    assert document.sections[0].page_height.inches == 11
    assert document.core_properties.author == ''
    assert next(p for p in document.paragraphs if p.text == 'Built beds').style.name == 'List Bullet'
    assert next(p for p in document.paragraphs if p.text == 'Plain point').style.name == 'Normal'


def test_word_export_all_default_fields_and_a4(sample_resume):
    document = Document(BytesIO(render_resume_docx(ResumeData.model_validate(sample_resume), page_size='A4', margins=(15, 16, 17, 18))))
    assert round(document.sections[0].page_width.mm) == 210
    assert round(document.sections[0].page_height.mm) == 297
    assert round(document.sections[0].left_margin.mm) == 17
    text = '\n'.join(p.text for p in document.paragraphs)
    for expected in ['JANE DOE', 'Acme Corp', 'MIT', 'OpenAPI Generator', 'PostgreSQL', 'Spanish (Conversational)', 'AWS Solutions Architect Associate', 'Employee of the Year 2022']:
        assert expected in text


@pytest.mark.integration
async def test_word_endpoint_session_owner_isolation_and_not_ready(isolated_backend_state, sample_resume, monkeypatch):
    monkeypatch.setattr(settings, 'auth_required', True)
    monkeypatch.setattr(settings, 'cookie_secure', False)
    database = isolated_backend_state
    with database._sync_write_session() as session:
        session.add_all([User(user_id='member-a', email='a@example.test', password_hash=hash_password('test-password-123')), User(user_id='member-b', email='b@example.test', password_hash=hash_password('test-password-123'))])
        session.commit()
    token = current_user_id.set('member-a')
    own = await database.create_resume('synthetic', processing_status='ready', processed_data=sample_resume)
    pending = await database.create_resume('synthetic pending', processing_status='pending')
    current_user_id.reset(token)
    token = current_user_id.set('member-b')
    other = await database.create_resume('other private content', processing_status='ready', processed_data=sample_resume)
    current_user_id.reset(token)
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
        assert (await client.get(f"/api/v1/resumes/{own['resume_id']}/docx")).status_code == 401
        signed_in = await client.post('/api/v1/auth/login', json={'email':'a@example.test','password':'test-password-123'}, headers={'Origin':'http://localhost:3000'})
        assert signed_in.status_code == 200
        exported = await client.get(f"/api/v1/resumes/{own['resume_id']}/docx")
        assert exported.status_code == 200
        assert exported.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        assert 'resume.docx' in exported.headers['content-disposition']
        assert exported.headers['cache-control'] == 'private, no-store'
        assert Document(BytesIO(exported.content)).paragraphs[0].text == 'JANE DOE'
        assert (await client.get(f"/api/v1/resumes/{other['resume_id']}/docx")).status_code == 404
        assert (await client.get(f"/api/v1/resumes/{pending['resume_id']}/docx")).status_code == 409
        assert (await client.get(f"/api/v1/resumes/{own['resume_id']}/docx?pageSize=INVALID")).status_code == 422


def test_word_export_matches_preview_geometry_typography_and_entry_layout(sample_resume):
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    document = Document(BytesIO(render_resume_docx(ResumeData.model_validate(sample_resume))))
    assert round(document.sections[0].page_width.mm) == 210
    assert round(document.sections[0].left_margin.mm) == 10
    assert document.styles['Title'].font.name == 'Georgia'
    assert document.styles['Title'].font.size.pt == 21
    assert document.styles['Normal'].font.name == 'Segoe UI'
    assert document.styles['Normal'].font.size.pt == 10.5
    assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    heading = next(p for p in document.paragraphs if p.text == 'EXPERIENCE')
    assert heading._p.xpath('./w:pPr/w:pBdr/w:bottom')
    job = next(p for p in document.paragraphs if p.text.startswith('Software Engineer'))
    assert '\t' in job.text
    assert job.paragraph_format.tab_stops[0].alignment == WD_TAB_ALIGNMENT.RIGHT
    assert 'Acme Corp' not in job.text
    company = next(p for p in document.paragraphs if p.text.startswith('Acme Corp'))
    assert company is not None
    tuned = Document(BytesIO(render_resume_docx(ResumeData.model_validate(sample_resume), font_size=5, header_font='mono', body_font='serif', compact=True)))
    assert tuned.styles['Normal'].font.size.pt == 12
    assert tuned.styles['Normal'].font.name == 'Georgia'
    assert tuned.styles['Title'].font.name == 'Consolas'
    assert tuned.styles['Heading 1'].paragraph_format.space_before < document.styles['Heading 1'].paragraph_format.space_before


def test_word_entry_spacing_and_skill_alignment_do_not_collapse(sample_resume):
    document = Document(BytesIO(render_resume_docx(ResumeData.model_validate(sample_resume))))
    first = next(p for p in document.paragraphs if p.text.startswith('Senior Backend Engineer\t'))
    second = next(p for p in document.paragraphs if p.text.startswith('Software Engineer\t'))
    bullet = next(p for p in document.paragraphs if p.text.startswith('Built REST APIs'))
    assert first.paragraph_format.space_before.pt == 0
    # Separate jobs retain a visible gap; wrapped bullets use the body line box.
    assert second.paragraph_format.space_before.pt > bullet.paragraph_format.space_after.pt
    assert bullet.paragraph_format.line_spacing.pt == pytest.approx(10.5 * 1.35, abs=.05)
    company = next(p for p in document.paragraphs if p.text.startswith('Acme Corp'))
    assert company.runs[0].font.name == 'Segoe UI Semibold'
    skills = next(p for p in document.paragraphs if p.text.startswith('Technical Skills:'))
    assert '\t' in skills.text
    assert skills.paragraph_format.left_indent.pt == 96
    assert skills.paragraph_format.first_line_indent.pt == -96
    assert 'PostgreSQL' in skills.text
